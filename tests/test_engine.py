"""Testes do motor de documentos (serviço, não HTTP)."""

from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx import Document

from app.config import settings
from app.docx_engine.writer import is_valid_docx
from app.models.blocks import HeadingContent, ParagraphContent, TableContent
from app.models.commands import Command, CommandBatch, FindRequest
from app.models.document import CreateDocumentRequest, OpenDocumentRequest, SaveDocumentRequest
from app.services.document_service import DocumentService
from app.storage.document_store import DocumentStore


def _service() -> DocumentService:
    return DocumentService(DocumentStore())


def test_create_document():
    service = _service()
    info = service.create(CreateDocumentRequest(filename="novo.docx", title="Relatório", author="IA"))
    assert info.id.startswith("doc_")
    assert info.filename == "novo.docx"
    assert info.title == "Relatório"


def test_add_and_update_and_delete_paragraph():
    service = _service()
    info = service.create(CreateDocumentRequest(filename="edit.docx"))
    pid = service.add_paragraph(info.id, ParagraphContent(text="Texto original."))
    assert pid.startswith("paragraph_")
    service.update_block(info.id, pid, {"text": "Texto alterado."})
    detail = service.get_block(info.id, pid)
    assert detail.text == "Texto alterado."
    service.delete_block(info.id, pid)
    structure = service.structure(info.id)
    assert all(block.id != pid for block in structure.blocks)


def test_add_heading_and_paragraph_after():
    service = _service()
    info = service.create(CreateDocumentRequest(filename="head.docx", title="Doc"))
    hid = service.add_heading(info.id, HeadingContent(text="Introdução", level=1))
    pid = service.add_paragraph(info.id, ParagraphContent(text="Texto introdutório."), after_id=hid)
    blocks = service.structure(info.id).blocks
    ids = [b.id for b in blocks]
    assert ids.index(hid) < ids.index(pid)
    heading = service.get_block(info.id, hid)
    assert heading.type.value == "heading"
    assert heading.level == 1
    assert heading.style == "Heading 1"


def test_add_table_and_set_cell():
    service = _service()
    info = service.create(CreateDocumentRequest(filename="tab.docx"))
    tid = service.add_table(
        info.id,
        TableContent(rows=2, columns=3, data=[["A", "B", "C"], ["1", "2", "3"]]),
    )
    service.set_cell(info.id, tid, 1, 2, "R$ 1.500,00")
    detail = service.get_block(info.id, tid)
    assert detail.rows == 2
    assert detail.columns == 3
    assert detail.cells[1][2] == "R$ 1.500,00"


def test_placeholder_replace_and_variables():
    service = _service()
    info = service.create(CreateDocumentRequest(filename="tpl.docx"))
    service.add_paragraph(info.id, ParagraphContent(text="Cliente: {{CLIENTE}} em {{DATA}}."))
    count = service.set_variables(info.id, {"CLIENTE": "Empresa ABC Ltda.", "DATA": "20/08/2026"})
    assert count >= 2
    matches = service.find(info.id, FindRequest(text="Empresa ABC"))
    assert matches
    assert "Empresa ABC Ltda." in matches[0].text


def test_replace_text_across_runs():
    service = _service()
    info = service.create(CreateDocumentRequest(filename="runs.docx"))
    service.add_paragraph(
        info.id,
        ParagraphContent(
            runs=[
                {"text": "{{CLI"},
                {"text": "ENTE}}", "bold": True},
            ]
        ),
    )
    replaced = service.replace_text(info.id, "{{CLIENTE}}", "Empresa ABC Ltda.")
    assert replaced >= 1
    structure = service.structure(info.id)
    para = next(b for b in structure.blocks if b.type.value == "paragraph" and b.text)
    assert "Empresa ABC Ltda." in (para.text or "")


def test_save_as_does_not_overwrite_original_and_reopen(tmp_path: Path):
    service = _service()
    original_name = "contrato.docx"
    info = service.create(CreateDocumentRequest(filename=original_name, title="Contrato"))
    service.add_paragraph(info.id, ParagraphContent(text="Cláusula original {{CLIENTE}}"))
    saved = service.save(info.id, SaveDocumentRequest(filename=original_name, keep_version=True))
    original = settings.documents_dir / original_name
    assert original.is_file()
    assert is_valid_docx(original)

    service.add_paragraph(info.id, ParagraphContent(text="Cláusula nova."))
    service.replace_text(info.id, "{{CLIENTE}}", "Empresa ABC")
    edited = service.save_as(info.id, "contrato_editado.docx")
    edited_path = settings.documents_dir / "contrato_editado.docx"
    assert edited_path.is_file()
    assert is_valid_docx(edited_path)
    with ZipFile(edited_path) as zf:
        assert "word/document.xml" in zf.namelist()

    reopened = service.open(OpenDocumentRequest(filename="contrato_editado.docx"))
    structure = service.structure(reopened.id)
    texts = " ".join(b.text or "" for b in structure.blocks)
    assert "Empresa ABC" in texts
    assert "Cláusula nova." in texts
    assert any(b.id.startswith("paragraph_") for b in structure.blocks)

    original_doc = Document(str(original))
    original_text = "\n".join(p.text for p in original_doc.paragraphs)
    assert "{{CLIENTE}}" in original_text or "Cláusula original" in original_text


def test_batch_commands_rollback_on_error():
    service = _service()
    info = service.create(CreateDocumentRequest(filename="batch.docx"))
    hid = service.add_heading(info.id, HeadingContent(text="Conclusão", level=1))
    before_count = len(service.structure(info.id).blocks)
    try:
        service.execute_commands(
            info.id,
            CommandBatch(
                commands=[
                    Command(action="add_paragraph", after_id=hid, text="Documento gerado automaticamente."),
                    Command(action="update_paragraph", target_id="paragraph_999", text="falha"),
                ]
            ),
        )
        assert False, "should have failed"
    except Exception:
        pass
    after_count = len(service.structure(info.id).blocks)
    assert after_count == before_count


def test_open_existing_extracts_structure(templates_dir: Path):
    path = settings.templates_dir / "proposta.docx"
    if not path.is_file():
        doc = Document()
        doc.add_heading("Proposta", 0)
        doc.add_heading("Introdução", 1)
        doc.add_paragraph("Cliente: {{CLIENTE}}")
        doc.add_table(rows=2, cols=2)
        doc.save(str(path))
    service = _service()
    info = service.open(OpenDocumentRequest(filename="proposta.docx", source="templates"))
    structure = service.structure(info.id)
    types = {b.type.value for b in structure.blocks}
    assert "heading" in types
    assert "paragraph" in types
    assert "table" in types
    assert all(b.id for b in structure.blocks)


def test_table_width_uses_twips_not_emu():
    from docx.oxml.ns import qn

    from app.docx_engine.xml_utils import cm_to_dxa

    service = _service()
    info = service.create(CreateDocumentRequest(filename="largura.docx"))
    tid = service.add_table(
        info.id,
        TableContent(rows=2, columns=2, data=[["A", "B"], ["1", "2"]], width_cm=15.5),
    )
    table = service.store.get(info.id).document.tables[0]
    tcw = table.cell(0, 0)._tc.find(qn("w:tcPr")).find(qn("w:tcW"))
    width = int(tcw.get(qn("w:w")))
    expected_col = cm_to_dxa(15.5) // 2
    assert abs(width - expected_col) < 5
    assert width < 20000  # twips (~15cm), never EMU (~millions)


def test_inspect_and_clone_existing_as_template(tmp_path: Path):
    service = _service()
    base = service.create(CreateDocumentRequest(filename="base_modelo.docx", title="Modelo"))
    service.add_heading(base.id, HeadingContent(text="Capa do modelo", level=1))
    service.add_paragraph(base.id, ParagraphContent(text="Cliente: {{CLIENTE}}"))
    service.add_table(base.id, TableContent(rows=2, columns=2, data=[["Col1", "Col2"], ["x", "y"]]))
    service.save(base.id, SaveDocumentRequest(filename="base_modelo.docx"))

    clone = service.create(
        CreateDocumentRequest(
            filename="a_partir_do_modelo.docx",
            template="base_modelo.docx",
            template_source="documents",
        )
    )
    inspected = service.inspect(clone.id)
    assert "{{CLIENTE}}" in inspected.placeholders or "CLIENTE" in inspected.placeholders
    assert inspected.outline
    assert inspected.tables
    assert inspected.language == "pt-BR"
    texts = " ".join(b.text or "" for b in service.structure(clone.id).blocks)
    assert "Capa do modelo" in texts


def test_toc_and_caption_fields():
    service = _service()
    info = service.create(CreateDocumentRequest(filename="sumario.docx"))
    hid = service.add_heading(info.id, HeadingContent(text="Introdução", level=1))
    service.execute_command(
        info.id,
        Command(action="insert_caption", after_id=hid, text="Fluxo de edição", kind="figura"),
    )
    result = service.execute_command(info.id, Command(action="insert_toc"))
    assert result["created_block_id"]
    xml = service.store.get(info.id).document.element.xml
    assert "TOC" in xml
    assert "SEQ Figura" in xml
