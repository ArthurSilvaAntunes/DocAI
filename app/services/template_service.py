"""Templates DOCX e substituição de placeholders {{VAR}}."""

from __future__ import annotations

import re

from docx.document import Document as DocxDocument

from app.docx_engine.search import replace_in_document

PLACEHOLDER_RE = re.compile(r"\{\{\s*([A-Za-z0-9_]+)\s*\}\}")


class TemplateService:
    def set_variables(self, document: DocxDocument, variables: dict[str, str]) -> int:
        total = 0
        for key, value in variables.items():
            total += replace_in_document(document, "{{" + key + "}}", value)
            total += replace_in_document(document, "{{ " + key + " }}", value)
        return total

    def list_placeholders(self, document: DocxDocument) -> list[str]:
        found: set[str] = set()
        for paragraph in document.paragraphs:
            for match in PLACEHOLDER_RE.findall(paragraph.text or ""):
                found.add(match)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for match in PLACEHOLDER_RE.findall(cell.text or ""):
                        found.add(match)
        return sorted(found)
