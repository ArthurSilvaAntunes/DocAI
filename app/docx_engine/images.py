"""Imagens inline no documento."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Emu
from docx.text.paragraph import Paragraph
from lxml import etree

from app.docx_engine import ids as id_mod
from app.docx_engine import xml_utils as xu
from app.docx_engine.fields import insert_caption, next_caption_number
from app.exceptions import AppError, NotFoundError


def add_image(
    document: DocxDocument,
    registry: id_mod.IdRegistry,
    image_bytes: bytes,
    *,
    filename: str = "image.png",
    width_cm: float | None = None,
    height_cm: float | None = None,
    alignment: str | None = None,
    caption: str | None = None,
    after_el: etree._Element | None = None,
    before_el: etree._Element | None = None,
) -> str:
    block_id, paragraph = insert_paragraph(
        document,
        registry,
        after_el=after_el,
        before_el=before_el,
        kind="image",
    )
    run = paragraph.add_run()
    kwargs: dict = {}
    if width_cm:
        kwargs["width"] = Cm(width_cm)
    if height_cm:
        kwargs["height"] = Cm(height_cm)
    try:
        run.add_picture(BytesIO(image_bytes), **kwargs)
    except Exception as exc:  # noqa: BLE001
        raise AppError("INVALID_IMAGE", f"Could not insert image '{filename}'.") from exc
    if alignment:
        paragraph.alignment = ALIGN_MAP.get(alignment, WD_ALIGN_PARAGRAPH.LEFT)
    if caption:
        insert_caption(
            document,
            registry,
            text=caption,
            kind="figura",
            after_el=paragraph._p,
            number=next_caption_number(document, "figura"),
        )
    return block_id


def replace_image(
    document: DocxDocument,
    block_id: str,
    image_bytes: bytes,
    *,
    width_cm: float | None = None,
    height_cm: float | None = None,
) -> None:
    paragraph = require_image_paragraph(document, block_id)
    drawings = list(paragraph._p.iter(qn("w:drawing")))
    if not drawings:
        raise AppError("IMAGE_NOT_FOUND", f"Block {block_id} does not contain an image.")
    for child in list(paragraph._p):
        if child.tag == xu.w("r"):
            paragraph._p.remove(child)
    run = paragraph.add_run()
    kwargs: dict = {}
    if width_cm:
        kwargs["width"] = Cm(width_cm)
    if height_cm:
        kwargs["height"] = Cm(height_cm)
    run.add_picture(BytesIO(image_bytes), **kwargs)


def require_image_paragraph(document: DocxDocument, block_id: str) -> Paragraph:
    el = id_mod.find_element(document, block_id)
    if el is None or el.tag != xu.w("p"):
        raise NotFoundError("BLOCK_NOT_FOUND", f"Block {block_id} was not found.")
    return wrap_paragraph(el, document)


def paragraph_has_image(p_el: etree._Element) -> bool:
    return p_el.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing") is not None


def _has_style(document: DocxDocument, name: str) -> bool:
    try:
        document.styles[name]
        return True
    except KeyError:
        return False


def load_image_bytes(path: Path) -> bytes:
    if not path.is_file():
        raise AppError("FILE_NOT_FOUND", f"Image was not found: {path.name}", status=404)
    return path.read_bytes()
