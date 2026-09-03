"""Cabeçalhos, rodapés e numeração de páginas com campos PAGE/NUMPAGES."""

from __future__ import annotations

from io import BytesIO

from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm

from app.docx_engine import xml_utils as xu
from app.models.document import HeaderFooterRequest


def _clear_container(container) -> None:
    for paragraph in list(container.paragraphs):
        p = paragraph._p
        parent = p.getparent()
        if parent is not None and len(list(parent.findall(qn("w:p")))) > 1:
            parent.remove(p)
        else:
            for child in list(p):
                if child.tag == xu.w("r"):
                    p.remove(child)


def set_header_footer(
    document: DocxDocument,
    payload: HeaderFooterRequest,
    *,
    kind: str = "header",
    document_name: str | None = None,
    image_bytes: bytes | None = None,
) -> None:
    for section in document.sections:
        container = section.header if kind == "header" else section.footer
        _clear_container(container)
        paragraph = container.paragraphs[0] if container.paragraphs else container.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        only_page = bool(payload.include_page_number) and bool(payload.number_only)
        if only_page:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if image_bytes:
            run = paragraph.add_run()
            run.add_picture(BytesIO(image_bytes), width=Cm(2.5))
            paragraph = container.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        parts: list[str] = []
        if payload.text:
            parts.append(payload.text)
        if payload.include_document_name and document_name:
            parts.append(document_name)
        if payload.include_date:
            from datetime import date

            parts.append(date.today().strftime("%d/%m/%Y"))
        if parts:
            paragraph.add_run("  |  ".join(parts) + ("  " if payload.include_page_number else ""))
        if payload.include_page_number:
            if only_page:
                xu.add_page_field(paragraph._p, label_before="", label_mid=None)
            else:
                xu.add_page_field(paragraph._p)
