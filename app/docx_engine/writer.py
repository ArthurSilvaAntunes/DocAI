"""Criação, gravação e clonagem de documentos DOCX sem recriar XML desnecessário."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, is_zipfile

from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.coreprops import CoreProperties

from app.docx_engine.language import DEFAULT_LANGUAGE, set_document_language
from app.docx_engine.sections import apply_page_setup
from app.exceptions import AppError
from app.models.document import PageSetup


def create_document(
    *,
    title: str | None = None,
    author: str | None = None,
    page: PageSetup | None = None,
    language: str | None = DEFAULT_LANGUAGE,
) -> DocxDocument:
    document = Document()
    if title:
        document.core_properties.title = title
    if author:
        document.core_properties.author = author
    apply_page_setup(document, page or PageSetup())
    if language:
        set_document_language(document, language)
    return document


def open_document(path: Path) -> DocxDocument:
    if not path.is_file():
        raise AppError("FILE_NOT_FOUND", f"File was not found: {path.name}", status=404)
    if not is_valid_docx(path):
        raise AppError("INVALID_DOCX", f"File '{path.name}' is not a valid DOCX archive.")
    try:
        return Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise AppError("INVALID_DOCX", f"Could not open '{path.name}' as a DOCX document.") from exc


def open_bytes(data: bytes) -> DocxDocument:
    if not is_valid_docx_bytes(data):
        raise AppError("INVALID_DOCX", "Uploaded file is not a valid DOCX archive.")
    try:
        return Document(BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise AppError("INVALID_DOCX", "Could not open uploaded file as a DOCX document.") from exc


def save_document(document: DocxDocument, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    if not is_valid_docx(path):
        raise AppError("CORRUPT_DOCX", "Saved file is not a valid DOCX archive.")


def document_bytes(document: DocxDocument) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def clone_from_bytes(data: bytes) -> DocxDocument:
    return open_bytes(data)


def is_valid_docx(path: Path) -> bool:
    if not is_zipfile(path):
        return False
    with ZipFile(path) as zf:
        names = set(zf.namelist())
    return "word/document.xml" in names and "[Content_Types].xml" in names


def is_valid_docx_bytes(data: bytes) -> bool:
    buffer = BytesIO(data)
    if not is_zipfile(buffer):
        return False
    buffer.seek(0)
    with ZipFile(buffer) as zf:
        names = set(zf.namelist())
    return "word/document.xml" in names and "[Content_Types].xml" in names


def core_title(document: DocxDocument) -> str | None:
    props: CoreProperties = document.core_properties
    return props.title or None


def core_author(document: DocxDocument) -> str | None:
    return document.core_properties.author or None
