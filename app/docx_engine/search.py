"""Busca e substituição, inclusive quando o texto está dividido entre runs."""

from __future__ import annotations

from docx.document import Document as DocxDocument
from docx.text.paragraph import Paragraph
from lxml import etree

from app.docx_engine import ids as id_mod
from app.docx_engine import xml_utils as xu
from app.docx_engine.paragraphs import wrap_paragraph
from app.models.blocks import BlockType
from app.models.responses import MatchItem


def paragraph_text(p_el: etree._Element) -> str:
    return "".join(t.text or "" for t in p_el.iter(xu.w("t")))


def replace_across_runs(paragraph: Paragraph, search: str, replace: str) -> int:
    if not search:
        return 0
    runs = paragraph.runs
    if not runs:
        return 0
    count = 0
    for run in runs:
        if search in (run.text or ""):
            n = run.text.count(search)
            run.text = run.text.replace(search, replace)
            count += n
    if count:
        return count

    full = "".join(r.text or "" for r in runs)
    if search not in full:
        return 0
    count = full.count(search)
    new_text = full.replace(search, replace)
    first = True
    for run in runs:
        if first:
            run.text = new_text
            first = False
        else:
            run.text = ""
    return count


def find_matches(
    document: DocxDocument,
    *,
    text: str | None = None,
    block_type: str | None = None,
    style: str | None = None,
    block_id: str | None = None,
) -> list[MatchItem]:
    from app.docx_engine.reader import classify_element, iter_blocks

    matches: list[MatchItem] = []
    needle = text.lower() if text else None
    for element, bid, kind in iter_blocks(document):
        if block_id and bid != block_id:
            continue
        if block_type and kind.value != block_type:
            continue
        if kind in (BlockType.paragraph, BlockType.heading, BlockType.list, BlockType.image):
            paragraph = wrap_paragraph(element, document)
            style_name = paragraph.style.name if paragraph.style is not None else ""
            if style and style_name != style:
                continue
            content = paragraph_text(element)
            if needle and needle not in content.lower():
                continue
            matches.append(MatchItem(block_id=bid, text=content[:300], type=kind.value))
            continue
        if kind == BlockType.table:
            from app.docx_engine.tables import read_cells, wrap_table

            table = wrap_table(element, document)
            joined = " ".join(cell for row in read_cells(table) for cell in row)
            if style:
                try:
                    if table.style and table.style.name != style:
                        continue
                except ValueError:
                    continue
            if needle and needle not in joined.lower():
                continue
            matches.append(MatchItem(block_id=bid, text=joined[:300], type=kind.value))
    return matches


def replace_in_document(document: DocxDocument, search: str, replace: str, block_id: str | None = None) -> int:
    total = 0
    if block_id:
        el = id_mod.find_element(document, block_id)
        if el is None:
            return 0
        if el.tag == xu.w("p"):
            return replace_across_runs(wrap_paragraph(el, document), search, replace)
        if el.tag == xu.w("tbl"):
            from app.docx_engine.tables import wrap_table

            table = wrap_table(el, document)
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        total += replace_across_runs(paragraph, search, replace)
            return total
    body = document.element.body
    for p_el in body.iter(xu.w("p")):
        total += replace_across_runs(wrap_paragraph(p_el, document), search, replace)
    return total
