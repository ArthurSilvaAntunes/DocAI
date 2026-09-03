"""Leitura de documentos existentes e extração da estrutura com IDs estáveis."""

from __future__ import annotations

from typing import Iterator

from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree

from app.docx_engine import ids as id_mod
from app.docx_engine import xml_utils as xu
from app.docx_engine.images import paragraph_has_image
from app.docx_engine.lists import list_type_of
from app.docx_engine.paragraphs import (
    heading_level_from_style,
    read_paragraph_properties,
    read_runs,
    wrap_paragraph,
)
from app.docx_engine.tables import read_cells, table_size, wrap_table
from app.exceptions import NotFoundError
from app.models.blocks import BlockDetail, BlockSummary, BlockType


def classify_element(element: etree._Element, document: DocxDocument) -> BlockType:
    if element.tag == xu.w("tbl"):
        return BlockType.table
    if xu.has_section_properties(element) and not _visible_text(element):
        return BlockType.section_break
    if xu.has_page_break(element) and not _visible_text(element):
        return BlockType.page_break
    if paragraph_has_image(element) and not _visible_text(element):
        return BlockType.image
    if xu.paragraph_has_num_pr(element):
        return BlockType.list
    paragraph = wrap_paragraph(element, document)
    style_name = paragraph.style.name if paragraph.style is not None else ""
    if heading_level_from_style(style_name) is not None:
        return BlockType.heading
    return BlockType.paragraph


def _visible_text(p_el: etree._Element) -> str:
    return "".join((t.text or "") for t in p_el.iter(xu.w("t"))).strip()


def kind_for_allocate(block_type: BlockType) -> str:
    mapping = {
        BlockType.heading: "heading",
        BlockType.paragraph: "paragraph",
        BlockType.table: "table",
        BlockType.image: "image",
        BlockType.list: "list",
        BlockType.page_break: "pagebreak",
        BlockType.section_break: "section",
    }
    return mapping[block_type]


def iter_blocks(document: DocxDocument) -> Iterator[tuple[etree._Element, str, BlockType]]:
    registry = id_mod.scan_registry(document)
    body = document.element.body
    for element in xu.iter_body_elements(body):
        block_type = classify_element(element, document)
        block_id = id_mod.ensure_block_id(element, kind_for_allocate(block_type), registry, document.element)
        yield element, block_id, block_type


def ensure_all_ids(document: DocxDocument) -> id_mod.IdRegistry:
    registry = id_mod.scan_registry(document)
    body = document.element.body
    for element in xu.iter_body_elements(body):
        block_type = classify_element(element, document)
        id_mod.ensure_block_id(element, kind_for_allocate(block_type), registry, document.element)
    return id_mod.scan_registry(document)


def get_structure(document: DocxDocument) -> list[BlockSummary]:
    blocks: list[BlockSummary] = []
    for element, block_id, block_type in iter_blocks(document):
        summary = BlockSummary(id=block_id, type=block_type)
        if block_type == BlockType.table:
            table = wrap_table(element, document)
            rows, cols = table_size(table)
            summary.rows = rows
            summary.columns = cols
            try:
                summary.style = table.style.name if table.style else None
            except ValueError:
                summary.style = None
        else:
            paragraph = wrap_paragraph(element, document)
            summary.text = _visible_text(element)
            summary.style = paragraph.style.name if paragraph.style is not None else None
            summary.level = heading_level_from_style(summary.style)
            if block_type == BlockType.list:
                summary.list_type = list_type_of(element)
                summary.level = xu.paragraph_num_level(element)
        blocks.append(summary)
    return blocks


def inspect_document(document: DocxDocument, *, document_id: str, filename: str) -> "DocumentInspect":
    from app.docx_engine.language import get_document_language
    from app.models.document import CaptionPreview, DocumentInspect, OutlineEntry, TablePreview
    from app.services.template_service import TemplateService

    outline: list[OutlineEntry] = []
    tables: list[TablePreview] = []
    images: list[str] = []
    captions: list[CaptionPreview] = []
    styles: set[str] = set()
    for element, block_id, block_type in iter_blocks(document):
        text = _visible_text(element)
        if block_type == BlockType.heading:
            paragraph = wrap_paragraph(element, document)
            style_name = paragraph.style.name if paragraph.style is not None else None
            if style_name:
                styles.add(style_name)
            outline.append(
                OutlineEntry(
                    id=block_id,
                    level=heading_level_from_style(style_name) or 1,
                    text=text,
                )
            )
        elif block_type == BlockType.table:
            table = wrap_table(element, document)
            rows, cols = table_size(table)
            header = read_cells(table)[0] if rows else []
            try:
                if table.style:
                    styles.add(table.style.name)
            except ValueError:
                pass
            tables.append(TablePreview(id=block_id, rows=rows, columns=cols, header=header))
        elif block_type == BlockType.image:
            images.append(block_id)
        else:
            paragraph = wrap_paragraph(element, document)
            style_name = paragraph.style.name if paragraph.style is not None else None
            if style_name:
                styles.add(style_name)
            lowered = text.lower()
            kind = None
            is_caption_style = (style_name or "") == "Caption"
            if is_caption_style:
                if lowered.startswith("figura"):
                    kind = "figura"
                elif lowered.startswith("tabela"):
                    kind = "tabela"
                elif lowered.startswith("quadro"):
                    kind = "quadro"
                else:
                    kind = "figura"
            if kind:
                captions.append(CaptionPreview(id=block_id, kind=kind, text=text))
    return DocumentInspect(
        document_id=document_id,
        filename=filename,
        title=document.core_properties.title or None,
        language=get_document_language(document),
        placeholders=TemplateService().list_placeholders(document),
        outline=outline,
        tables=tables,
        images=images,
        captions=captions,
        styles_in_use=sorted(styles),
        block_count=len(get_structure(document)),
    )


def get_block_detail(document: DocxDocument, block_id: str) -> BlockDetail:
    element = id_mod.find_element(document, block_id)
    if element is None:
        raise NotFoundError("BLOCK_NOT_FOUND", f"Block {block_id} was not found.")
    block_type = classify_element(element, document)
    detail = BlockDetail(id=block_id, type=block_type)
    if block_type == BlockType.table:
        table = wrap_table(element, document)
        rows, cols = table_size(table)
        detail.rows = rows
        detail.columns = cols
        detail.cells = read_cells(table)
        try:
            detail.style = table.style.name if table.style else None
        except ValueError:
            detail.style = None
        return detail
    paragraph = wrap_paragraph(element, document)
    detail.text = _visible_text(element)
    detail.style = paragraph.style.name if paragraph.style is not None else None
    detail.level = heading_level_from_style(detail.style)
    detail.properties = read_paragraph_properties(paragraph)
    detail.runs = read_runs(paragraph)
    if block_type == BlockType.list:
        detail.list_type = list_type_of(element)
        detail.level = xu.paragraph_num_level(element)
    if block_type == BlockType.image:
        detail.extra = {"has_image": True}
    return detail


def resolve_anchor(document: DocxDocument, block_id: str) -> etree._Element:
    element = id_mod.find_element(document, block_id)
    if element is None:
        raise NotFoundError("BLOCK_NOT_FOUND", f"Block {block_id} was not found.")
    return element
