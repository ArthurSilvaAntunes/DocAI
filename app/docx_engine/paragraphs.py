"""Operações de parágrafos, headings e runs."""

from __future__ import annotations

from typing import Optional

from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_UNDERLINE
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree

from app.docx_engine import ids as id_mod
from app.docx_engine import xml_utils as xu
from app.exceptions import AppError, NotFoundError
from app.models.blocks import Alignment, ParagraphProperties, RunFormat, RunModel

ALIGN_MAP = {
    Alignment.left: WD_ALIGN_PARAGRAPH.LEFT,
    Alignment.center: WD_ALIGN_PARAGRAPH.CENTER,
    Alignment.right: WD_ALIGN_PARAGRAPH.RIGHT,
    Alignment.justify: WD_ALIGN_PARAGRAPH.JUSTIFY,
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

HIGHLIGHT_MAP = {
    "yellow": "yellow",
    "green": "green",
    "cyan": "cyan",
    "magenta": "magenta",
    "blue": "blue",
    "red": "red",
    "darkBlue": "darkBlue",
    "darkCyan": "darkCyan",
    "darkGreen": "darkGreen",
    "darkMagenta": "darkMagenta",
    "darkRed": "darkRed",
    "darkYellow": "darkYellow",
    "darkGray": "darkGray",
    "lightGray": "lightGray",
    "black": "black",
    "white": "white",
}


def wrap_paragraph(element: etree._Element, document: DocxDocument) -> Paragraph:
    return Paragraph(element, document)


def heading_style(level: int) -> str:
    level = max(1, min(9, int(level)))
    return f"Heading {level}"


def heading_level_from_style(name: str | None) -> Optional[int]:
    if not name:
        return None
    if name.startswith("Heading "):
        try:
            return int(name.split(" ", 1)[1])
        except ValueError:
            return None
    return None


def apply_paragraph_properties(paragraph: Paragraph, props: ParagraphProperties | dict | None) -> None:
    if not props:
        return
    data = props.model_dump(exclude_none=True) if isinstance(props, ParagraphProperties) else {k: v for k, v in props.items() if v is not None}
    style = data.get("style")
    if style:
        try:
            paragraph.style = style
        except (KeyError, ValueError) as exc:
            raise AppError("STYLE_NOT_FOUND", f"Style '{style}' was not found.") from exc
    alignment = data.get("alignment")
    if alignment:
        paragraph.alignment = ALIGN_MAP.get(alignment, WD_ALIGN_PARAGRAPH.LEFT)
    fmt = paragraph.paragraph_format
    if "indent_left_cm" in data:
        fmt.left_indent = Cm(float(data["indent_left_cm"]))
    if "first_line_indent_cm" in data:
        fmt.first_line_indent = Cm(float(data["first_line_indent_cm"]))
    if "space_before_pt" in data:
        fmt.space_before = Pt(float(data["space_before_pt"]))
    if "space_after_pt" in data:
        fmt.space_after = Pt(float(data["space_after_pt"]))
    if "line_spacing" in data:
        fmt.line_spacing = float(data["line_spacing"])
    if "keep_with_next" in data:
        fmt.keep_with_next = bool(data["keep_with_next"])
    if "page_break_before" in data:
        fmt.page_break_before = bool(data["page_break_before"])


def apply_run_format(run: Run, fmt: RunFormat | dict | None) -> None:
    if not fmt:
        return
    data = fmt.model_dump(exclude_none=True) if isinstance(fmt, RunFormat) else {k: v for k, v in fmt.items() if v is not None}
    if "bold" in data:
        run.bold = bool(data["bold"])
    if "italic" in data:
        run.italic = bool(data["italic"])
    if "underline" in data:
        run.underline = WD_UNDERLINE.SINGLE if data["underline"] else WD_UNDERLINE.NONE
    if "strikethrough" in data:
        run.font.strike = bool(data["strikethrough"])
    if "font" in data:
        run.font.name = str(data["font"])
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = xu.make("rFonts")
            rpr.append(rfonts)
        for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
            rfonts.set(qn(f"w:{attr}"), str(data["font"]))
    if "font_size" in data:
        run.font.size = Pt(float(data["font_size"]))
    if "color" in data:
        run.font.color.rgb = _parse_color(str(data["color"]))
    if "superscript" in data:
        run.font.superscript = bool(data["superscript"])
    if "subscript" in data:
        run.font.subscript = bool(data["subscript"])
    if "highlight" in data:
        highlight = HIGHLIGHT_MAP.get(str(data["highlight"]), str(data["highlight"]))
        rpr = run._element.get_or_add_rPr()
        xu.remove_children(rpr, "highlight")
        rpr.append(xu.make("highlight", val=highlight))


def _parse_color(value: str) -> RGBColor:
    raw = value.lstrip("#")
    if len(raw) != 6:
        raise AppError("INVALID_COLOR", f"Color '{value}' must be RRGGBB.")
    return RGBColor(int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16))


def set_runs(paragraph: Paragraph, runs: list[RunModel] | list[dict] | None, text: str | None = None) -> None:
    _clear_runs(paragraph)
    if runs:
        for item in runs:
            model = item if isinstance(item, RunModel) else RunModel.model_validate(item)
            run = paragraph.add_run(model.text)
            apply_run_format(run, model)
        return
    if text is not None:
        paragraph.add_run(text)


def _clear_runs(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag == xu.w("r"):
            paragraph._p.remove(child)


def read_runs(paragraph: Paragraph) -> list[RunModel]:
    result: list[RunModel] = []
    for run in paragraph.runs:
        highlight = None
        rpr = run._element.find(qn("w:rPr"))
        if rpr is not None:
            hl = rpr.find(qn("w:highlight"))
            if hl is not None:
                highlight = hl.get(qn("w:val"))
        color = None
        if run.font.color and run.font.color.rgb:
            color = str(run.font.color.rgb)
        result.append(
            RunModel(
                text=run.text or "",
                bold=bool(run.bold) if run.bold is not None else None,
                italic=bool(run.italic) if run.italic is not None else None,
                underline=bool(run.underline) if run.underline else None,
                strikethrough=bool(run.font.strike) if run.font.strike is not None else None,
                font=run.font.name,
                font_size=run.font.size.pt if run.font.size else None,
                highlight=highlight,
                color=color,
                superscript=bool(run.font.superscript) if run.font.superscript else None,
                subscript=bool(run.font.subscript) if run.font.subscript else None,
            )
        )
    return result


def read_paragraph_properties(paragraph: Paragraph) -> ParagraphProperties:
    fmt = paragraph.paragraph_format
    alignment = None
    if paragraph.alignment is not None:
        reverse = {
            WD_ALIGN_PARAGRAPH.LEFT: Alignment.left,
            WD_ALIGN_PARAGRAPH.CENTER: Alignment.center,
            WD_ALIGN_PARAGRAPH.RIGHT: Alignment.right,
            WD_ALIGN_PARAGRAPH.JUSTIFY: Alignment.justify,
        }
        alignment = reverse.get(paragraph.alignment)
    style_name = paragraph.style.name if paragraph.style is not None else None
    return ParagraphProperties(
        style=style_name,
        alignment=alignment,
        indent_left_cm=fmt.left_indent.cm if fmt.left_indent else None,
        first_line_indent_cm=fmt.first_line_indent.cm if fmt.first_line_indent else None,
        space_before_pt=fmt.space_before.pt if fmt.space_before else None,
        space_after_pt=fmt.space_after.pt if fmt.space_after else None,
        line_spacing=float(fmt.line_spacing) if isinstance(fmt.line_spacing, (int, float)) else None,
        keep_with_next=fmt.keep_with_next,
        page_break_before=fmt.page_break_before,
    )


def new_paragraph_element() -> etree._Element:
    return xu.make("p")


def insert_paragraph(
    document: DocxDocument,
    registry: id_mod.IdRegistry,
    *,
    text: str | None = None,
    runs: list | None = None,
    properties: ParagraphProperties | dict | None = None,
    after_el: etree._Element | None = None,
    before_el: etree._Element | None = None,
    kind: str = "paragraph",
    style: str | None = None,
) -> tuple[str, Paragraph]:
    p_el = new_paragraph_element()
    body = document.element.body
    if before_el is not None:
        before_el.addprevious(p_el)
    elif after_el is not None:
        after_el.addnext(p_el)
    else:
        sect = body.find(xu.w("sectPr"))
        if sect is not None:
            sect.addprevious(p_el)
        else:
            body.append(p_el)
    paragraph = wrap_paragraph(p_el, document)
    if style:
        apply_paragraph_properties(paragraph, {"style": style})
    apply_paragraph_properties(paragraph, properties)
    set_runs(paragraph, runs, text)
    block_id = id_mod.ensure_block_id(p_el, kind, registry, document.element)
    return block_id, paragraph


def update_paragraph(
    paragraph: Paragraph,
    *,
    text: str | None = None,
    runs: list | None = None,
    properties: ParagraphProperties | dict | None = None,
) -> None:
    apply_paragraph_properties(paragraph, properties)
    if runs is not None or text is not None:
        set_runs(paragraph, runs, text)


def add_page_break_paragraph(document: DocxDocument, registry: id_mod.IdRegistry, after_el: etree._Element | None = None) -> str:
    block_id, paragraph = insert_paragraph(document, registry, after_el=after_el, kind="pagebreak")
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    return block_id


def require_paragraph(document: DocxDocument, block_id: str) -> Paragraph:
    el = id_mod.find_element(document, block_id)
    if el is None or el.tag != xu.w("p"):
        raise NotFoundError("BLOCK_NOT_FOUND", f"Block {block_id} was not found.")
    return wrap_paragraph(el, document)
