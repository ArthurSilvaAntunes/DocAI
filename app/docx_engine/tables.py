"""Operações de tabelas no DOCX."""

from __future__ import annotations

from copy import deepcopy

from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.table import Table, _Cell
from lxml import etree

from app.docx_engine import ids as id_mod
from app.docx_engine import xml_utils as xu
from app.docx_engine.language import DEFAULT_LANGUAGE, apply_run_language
from app.docx_engine.paragraphs import ALIGN_MAP
from app.exceptions import AppError, NotFoundError

# A4 com margens de 2,54 cm cabe ~15,9 cm. 16 cm estoura a página.
MAX_TABLE_WIDTH_CM = 15.5


def wrap_table(element: etree._Element, document: DocxDocument) -> Table:
    return Table(element, document)


def add_table(
    document: DocxDocument,
    registry: id_mod.IdRegistry,
    rows: int,
    cols: int,
    data: list[list[str]] | None = None,
    style: str | None = None,
    after_el: etree._Element | None = None,
    before_el: etree._Element | None = None,
    width_cm: float | None = None,
) -> tuple[str, Table]:
    table = document.add_table(rows=rows, cols=cols)
    tbl = table._tbl
    if before_el is not None:
        before_el.addprevious(tbl)
    elif after_el is not None:
        after_el.addnext(tbl)
    if style:
        try:
            table.style = style
        except (KeyError, ValueError):
            table.style = "Table Grid"
    else:
        try:
            table.style = "Table Grid"
        except (KeyError, ValueError):
            xu.set_table_borders(tbl)
    apply_table_layout(table, width_cm=width_cm)
    if table.rows:
        xu.mark_header_row(table.rows[0]._tr)
    if data:
        for r, row_data in enumerate(data):
            if r >= rows:
                break
            for c, value in enumerate(row_data):
                if c >= cols:
                    break
                set_cell_text(table.cell(r, c), str(value))
    block_id = id_mod.ensure_block_id(tbl, "table", registry, document.element)
    return block_id, table


def require_table(document: DocxDocument, block_id: str) -> Table:
    el = id_mod.find_element(document, block_id)
    if el is None or el.tag != xu.w("tbl"):
        raise NotFoundError("BLOCK_NOT_FOUND", f"Block {block_id} was not found.")
    return wrap_table(el, document)


def set_cell_text(
    cell: _Cell,
    text: str,
    *,
    bold: bool | None = None,
    alignment: str | None = None,
    shading: str | None = None,
) -> None:
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    if paragraph.runs:
        paragraph.runs[0].text = text
        for extra in paragraph.runs[1:]:
            extra.text = ""
        if bold is not None:
            paragraph.runs[0].bold = bold
        apply_run_language(paragraph.runs[0]._element, DEFAULT_LANGUAGE)
    else:
        run = paragraph.add_run(text)
        if bold is not None:
            run.bold = bold
        apply_run_language(run._element, DEFAULT_LANGUAGE)
    for extra_p in cell.paragraphs[1:]:
        for run in extra_p.runs:
            run.text = ""
    if alignment:
        paragraph.alignment = ALIGN_MAP.get(alignment, WD_ALIGN_PARAGRAPH.LEFT)
    if shading:
        xu.set_cell_shading(cell._tc, shading)


def add_row(table: Table, values: list[str] | None = None) -> None:
    row = table.add_row()
    if values:
        for idx, value in enumerate(values):
            if idx < len(row.cells):
                set_cell_text(row.cells[idx], value)


def delete_row(table: Table, index: int) -> None:
    if index < 0 or index >= len(table.rows):
        raise AppError("ROW_NOT_FOUND", f"Row {index} was not found.")
    if len(table.rows) <= 1:
        raise AppError("INVALID_TABLE", "A table must keep at least one row.")
    tbl = table._tbl
    tr = table.rows[index]._tr
    tbl.remove(tr)


def add_column(table: Table, values: list[str] | None = None) -> None:
    tbl = table._tbl
    grid = xu.first_child(tbl, "tblGrid")
    if grid is not None:
        grid_cols = xu.children(grid, "gridCol")
        width = xu.get_attr(grid_cols[-1], "w") if grid_cols else "1200"
        grid.append(xu.make("gridCol", w=width or "1200"))
    for r_idx, row in enumerate(table.rows):
        last = row.cells[-1]._tc
        new_tc = deepcopy(last)
        for p in new_tc.iter(xu.w("t")):
            p.text = ""
        row._tr.append(new_tc)
        if values and r_idx < len(values):
            set_cell_text(row.cells[-1], values[r_idx])


def delete_column(table: Table, index: int) -> None:
    if index < 0 or index >= len(table.columns):
        raise AppError("COLUMN_NOT_FOUND", f"Column {index} was not found.")
    if len(table.columns) <= 1:
        raise AppError("INVALID_TABLE", "A table must keep at least one column.")
    tbl = table._tbl
    grid = xu.first_child(tbl, "tblGrid")
    if grid is not None:
        cols = xu.children(grid, "gridCol")
        if index < len(cols):
            grid.remove(cols[index])
    for row in table.rows:
        cells = row._tr.findall(xu.w("tc"))
        if index < len(cells):
            row._tr.remove(cells[index])


def merge_cells(table: Table, start_row: int, start_col: int, end_row: int, end_col: int) -> None:
    try:
        start = table.cell(start_row, start_col)
        end = table.cell(end_row, end_col)
    except IndexError as exc:
        raise AppError("CELL_NOT_FOUND", "Cell coordinates are out of range.") from exc
    start.merge(end)


def read_cells(table: Table) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in table.rows]


def table_size(table: Table) -> tuple[int, int]:
    rows = len(table.rows)
    cols = len(table.columns) if rows else 0
    return rows, cols


def apply_table_layout(table: Table, width_cm: float | None = None) -> None:
    rows, cols = table_size(table)
    if cols <= 0:
        return
    if width_cm:
        width = min(float(width_cm), MAX_TABLE_WIDTH_CM)
        xu.set_table_width_dxa(table._tbl, xu.cm_to_dxa(width), cols)
        return
    xu.set_table_full_width(table._tbl)


def normalize_table(table: Table, width_cm: float | None = None) -> None:
    """Repara tabelas com tcW em EMU (valor absurdo) e aplica largura da página."""
    apply_table_layout(table, width_cm=width_cm or MAX_TABLE_WIDTH_CM)
    if table.rows:
        xu.mark_header_row(table.rows[0]._tr)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    apply_run_language(run._element, DEFAULT_LANGUAGE)
