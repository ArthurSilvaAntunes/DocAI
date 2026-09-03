from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.config import settings


def build_proposta(path: Path | None = None) -> Path:
    settings.templates_dir.mkdir(parents=True, exist_ok=True)
    dest = path or (settings.templates_dir / "proposta.docx")
    doc = Document()
    title = doc.add_heading("Proposta Comercial", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Cliente: {{CLIENTE}}")
    doc.add_paragraph("Data: {{DATA}}")
    doc.add_paragraph("Valor: {{VALOR}}")
    doc.add_paragraph("Responsável: {{RESPONSAVEL}}")
    doc.add_heading("Escopo", 1)
    doc.add_paragraph("Descrição do escopo do trabalho a ser realizado.")
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Valor"
    table.cell(1, 0).text = "Serviço"
    table.cell(1, 1).text = "{{VALOR}}"
    doc.add_heading("Conclusão", 1)
    p = doc.add_paragraph("Documento gerado a partir do template de proposta.")
    p.runs[0].font.size = Pt(11)
    doc.save(str(dest))
    return dest


if __name__ == "__main__":
    print(build_proposta())
